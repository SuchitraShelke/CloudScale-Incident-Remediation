"""3-year TCO / ROI model: builds docs/financial/tco-roi-model.xlsx (live formulas) and prints the same
results computed independently in Python, so the two can be cross-checked.

    uv run --with openpyxl python scripts/tco_model.py            # build + print Python results
    uv run --with openpyxl --with formulas python scripts/tco_model.py --verify   # also evaluate the xlsx

Inputs are the blueprint's (CapstoneProjectPlan_v2.md §12, §17) except the Year-0 build line, which is
the actual solo build (1 developer + Claude, ~24 h) instead of 10 engineers x 12 h.
"""

import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parents[1] / "docs" / "financial" / "tco-roi-model.xlsx"

# ------------------------------------------------------------------ inputs (single source)
I = {
    "rate": 150, "inc_y1": 200, "growth": 0.30,
    "adoption": [0.70, 0.90, 0.95], "auto": [0.40, 0.50, 0.55], "hitl": [0.40, 0.35, 0.32],
    "min_auto": 38, "min_hitl": 25, "min_esc": 10, "bad_rate": [0.02, 0.02, 0.015], "bad_min": 120,
    "maint_h": [20, 20, 16], "llm_eval": 100, "discount": 0.10, "azure_dr": 500,
    "build_h": 24, "prod": 72000, "landing": 2000, "policy": 1500,
    "infra": [("EKS control plane", [73, 73, 73]), ("App node group (m6i.large x3/4/5)", [210, 280, 350]),
              ("Guard node (c6i.xlarge, Llama Guard)", [124, 124, 124]),
              ("RDS PostgreSQL (db.t4g.medium, Multi-AZ, 50 GB)", [105, 105, 105]),
              ("ElastiCache Redis (cache.t4g.medium)", [47, 47, 47]), ("Networking (ALB + NAT + egress)", [70, 70, 70]),
              ("Observability storage (30-day)", [40, 40, 40]), ("Logs, registry, secrets", [20, 20, 20])],
    "price": {"sonnet": (2.00, 10.00), "haiku": (1.00, 5.00)},
    "agents": [("Triage", 8192, 2048), ("Planner", 4096, 3072), ("Evaluator summary", 2048, 1024)],
    "p1_share": 0.40, "cache_hit": 0.20, "overhead": 1.3,
    "shift": 0.15, "slow_adoption": [0.50, 0.70, 0.75], "vol_cut": 0.30,
}


# ------------------------------------------------------------------ independent Python computation
def compute() -> dict:
    cost = {m: [(i * p[0] + o * p[1]) / 1e6 for _, i, o in I["agents"]] for m, p in I["price"].items()}
    s, h, p1 = cost["sonnet"], cost["haiku"], I["p1_share"]
    all_sonnet = sum(s)
    routed_p1, routed_p2 = s[0] + s[1] + h[2], sum(h)
    routed = p1 * routed_p1 + (1 - p1) * routed_p2
    tp_part = p1 * (s[0] + s[1]) + (1 - p1) * (h[0] + h[1])
    strategies = {"all_sonnet": all_sonnet * I["overhead"], "routed": routed * I["overhead"],
                  "routed_cache": (routed - I["cache_hit"] * tp_part) * I["overhead"]}
    cpi = strategies["routed_cache"]

    inc = [I["inc_y1"] * (1 + I["growth"]) ** n for n in range(3)]
    infra = [sum(v[y] for _, v in I["infra"]) for y in range(3)]
    llm_prod = [inc[y] * cpi for y in range(3)]
    monthly = [infra[y] + llm_prod[y] + I["llm_eval"] + I["maint_h"][y] * I["rate"] for y in range(3)]
    opex = [m * 12 for m in monthly]
    capex = I["build_h"] * I["rate"] + I["prod"] + I["landing"] + I["policy"]

    benefit, via = [], []
    for y in range(3):
        v = inc[y] * I["adoption"][y]
        a, hi = v * I["auto"][y], v * I["hitl"][y]
        e = v * (1 - I["auto"][y] - I["hitl"][y])
        minutes = a * I["min_auto"] + hi * I["min_hitl"] + e * I["min_esc"] - (a + hi) * I["bad_rate"][y] * I["bad_min"]
        via.append(v)
        benefit.append(minutes / 60 * 12 * I["rate"])

    nets = [-capex] + [benefit[y] - opex[y] for y in range(3)]
    cum = [sum(nets[: i + 1]) for i in range(4)]
    pv = [nets[i] / (1 + I["discount"]) ** i for i in range(4)]
    npv = sum(pv)
    total_cost, total_benefit = capex + sum(opex), sum(benefit)
    # Month of operation in which the cumulative cash flow turns positive (same rule as the xlsx ROUNDUP).
    if cum[1] >= 0:
        payback = math.ceil(-cum[0] / (nets[1] / 12))
    elif cum[2] >= 0:
        payback = 12 + math.ceil(-cum[1] / (nets[2] / 12))
    else:
        payback = 24 + math.ceil(-cum[2] / (nets[3] / 12))
    base = cum[3]
    llm_total = sum(llm_prod[y] + I["llm_eval"] for y in range(3)) * 12
    sens = {
        "Base": base,
        "Lower automation (15 pts automated -> HITL)": base - sum(
            via[y] * I["shift"] * (I["min_auto"] - I["min_hitl"]) * 12 / 60 * I["rate"] for y in range(3)),
        "Slower adoption (50/70/75 %)": base - sum(benefit[y] * (1 - I["slow_adoption"][y] / I["adoption"][y])
                                                  for y in range(3)),
        "Lower volume (-30 % incidents)": base - I["vol_cut"] * total_benefit + I["vol_cut"] * sum(llm_prod) * 12,
        "LLM prices x2": base - llm_total,
        "Productionization +50 %": base - 0.5 * I["prod"],
        "Azure warm DR (+$500/month)": base - I["azure_dr"] * 36,
    }
    strat_3yr = {k: sum(inc[y] * 12 * v for y in range(3)) for k, v in strategies.items()}
    return {"agent_cost_sonnet": s, "agent_cost_haiku": h, "strategies": strategies, "strategy_3yr": strat_3yr,
            "cost_per_incident": cpi, "p1_cpi": routed_p1 * I["overhead"], "incidents_pm": inc, "infra_pm": infra,
            "llm_prod_pm": llm_prod, "monthly": monthly, "opex": opex, "capex": capex, "benefit": benefit,
            "via": via, "nets": nets, "cum": cum, "pv": pv, "npv": npv, "total_cost": total_cost,
            "total_benefit": total_benefit, "bcr": total_benefit / total_cost, "roi": cum[3] / total_cost,
            "payback_month": payback, "llm_share": llm_total / total_cost, "sensitivity": sens,
            "sre_hours_y1": benefit[0] / I["rate"], "sre_hours_y3": benefit[2] / I["rate"]}


# ------------------------------------------------------------------ workbook with live formulas
FONT = "Arial"
BLUE, GREEN = Font(name=FONT, color="0000FF"), Font(name=FONT, color="008000")
BOLD, NORMAL = Font(name=FONT, bold=True), Font(name=FONT)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
USD, USD2, USD4 = '$#,##0;($#,##0);"-"', '$#,##0.00;($#,##0.00);"-"', '$#,##0.0000;($#,##0.0000);"-"'
PCT, NUM, NUM1 = '0.0%;(0.0%);"-"', '#,##0;(#,##0);"-"', '#,##0.0;(#,##0.0);"-"'
THIN = Border(bottom=Side(style="thin", color="999999"))


def header(ws, row, labels, widths=None):
    for c, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, THIN
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left", wrap_text=True)
    for c, w in enumerate(widths or [], start=1):
        ws.column_dimensions[chr(64 + c)].width = w


def put(ws, ref, value, fmt=None, font=None, note=None, fill=None):
    cell = ws[ref]
    cell.value = value
    is_formula = isinstance(value, str) and value.startswith("=")
    links_other_sheet = is_formula and "!" in value
    cell.font = font or (GREEN if links_other_sheet else NORMAL if is_formula else BLUE)
    if fmt:
        cell.number_format = fmt
    if note:
        cell.comment = Comment(note, "CloudScale model")
    if fill:
        cell.fill = fill
    return cell


def label(ws, ref, text, bold=False):
    ws[ref].value, ws[ref].font = text, BOLD if bold else NORMAL


def build() -> None:
    wb = Workbook()
    Y = ["Y1", "Y2", "Y3"]

    # ---------------- Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"].value, ws["A1"].font = "Inputs and assumptions", TITLE
    ws["A2"].value, ws["A2"].font = ("Blue = input you can change · black = formula · green = link to another sheet. "
                                     "Source: CapstoneProjectPlan_v2.md §17 unless noted."), Font(name=FONT, italic=True)
    header(ws, 3, ["Assumption", *Y, "Note"], [46, 12, 12, 12, 60])
    rows = [
        (4, "SRE fully loaded rate ($/h)", [I["rate"]], USD, "Blueprint §17.1"),
        (5, "Incidents per month, Y1", [I["inc_y1"]], NUM, "Blueprint §17.1"),
        (6, "Incident growth per year", [I["growth"]], PCT, "+30 % YoY"),
        (7, "Share of incidents routed through the platform", I["adoption"], PCT, "Adoption ramp; v1 assumed 100 % on day 1"),
        (8, "Outcome mix: automated", I["auto"], PCT, "Informed by the scenario gate mix"),
        (9, "Outcome mix: human-approved (HITL)", I["hitl"], PCT, ""),
        (11, "Minutes saved per automated incident", [I["min_auto"]], NUM, "45 min manual MTTR -> ~7 min incl. post-review"),
        (12, "Minutes saved per HITL incident", [I["min_hitl"]], NUM, "45 -> 20 min"),
        (13, "Minutes saved per escalated incident", [I["min_esc"]], NUM, "Diagnosis already done when a human picks it up"),
        (14, "Bad-remediation rate (automated + HITL)", I["bad_rate"], PCT, "Share that makes things worse"),
        (15, "Extra minutes per bad remediation", [I["bad_min"]], NUM, "Cost of cleaning up a bad fix"),
        (16, "Platform maintenance (hours / month)", I["maint_h"], NUM, "20 h x $150 = $3,000 / month"),
        (17, "LLM evaluation + staging ($ / month)", [I["llm_eval"]], USD, "Prompt iteration, eval runs"),
        (18, "Discount rate", [I["discount"]], PCT, ""),
        (19, "Azure warm DR ($ / month, sensitivity only)", [I["azure_dr"]], USD, "ADR-007"),
    ]
    for r, text, vals, fmt, note in rows:
        label(ws, f"A{r}", text)
        for i, v in enumerate(vals):
            put(ws, f"{'BCD'[i]}{r}", v, fmt)
        ws[f"E{r}"].value, ws[f"E{r}"].font = note, NORMAL
    label(ws, "A10", "Outcome mix: escalated")
    for col in "BCD":
        put(ws, f"{col}10", f"=1-{col}8-{col}9", PCT)
    ws["E10"].value = "Remainder of the mix"
    for r in (4, 5, 6):   # single-value rows: only column B is an input
        for col in "CD":
            ws[f"{col}{r}"].value = None

    ws["A21"].value, ws["A21"].font = "Year 0 (one-time)", BOLD
    capex_rows = [
        (22, "Capstone build (hours)", I["build_h"], NUM,
         "Actual: 1 developer + Claude, 2 days (~24 h). Blueprint assumed 10 engineers x 12 h = 120 h"),
        (23, "Productionization ($): real K8s/cloud adapters, OIDC, mesh mTLS, pen test", I["prod"], USD,
         "2 engineers x 6 weeks x $150/h"),
        (24, "Cloud landing zone + hardening ($)", I["landing"], USD, ""),
        (25, "Policy authoring + review: OPA, gate policy ($)", I["policy"], USD, ""),
    ]
    for r, text, v, fmt, note in capex_rows:
        label(ws, f"A{r}", text)
        put(ws, f"B{r}", v, fmt, fill=KEY_FILL if r == 22 else None)
        ws[f"E{r}"].value = note

    ws["A27"].value, ws["A27"].font = "Infrastructure ($ / month, AWS EKS us-east-1 on-demand estimates)", BOLD
    for k, (text, vals) in enumerate(I["infra"]):
        r = 28 + k
        label(ws, f"A{r}", text)
        for i, v in enumerate(vals):
            put(ws, f"{'BCD'[i]}{r}", v, USD)
    ws["E28"].value = "Validate with the AWS Pricing Calculator before external use (ADR-007)"

    ws["A38"].value, ws["A38"].font = "Sensitivity levers", BOLD
    label(ws, "A39", "Lower automation: points shifted automated -> HITL")
    put(ws, "B39", I["shift"], PCT)
    label(ws, "A40", "Slower adoption case")
    for i, v in enumerate(I["slow_adoption"]):
        put(ws, f"{'BCD'[i]}40", v, PCT)
    label(ws, "A41", "Lower volume case: incident cut")
    put(ws, "B41", I["vol_cut"], PCT)

    # ---------------- Token Economics
    te = wb.create_sheet("Token Economics")
    te["A1"].value, te["A1"].font = "Token economics per incident", TITLE
    te["A2"].value = "Anthropic first-party list prices per 1M tokens (verify before use). Thinking tokens bill as output."
    header(te, 4, ["Model", "Input $/1M", "Output $/1M"], [34, 14, 14, 16, 16, 16])
    for r, (name, (pi, po)) in zip((5, 6), (("claude-sonnet-5", I["price"]["sonnet"]),
                                            ("claude-haiku-4-5", I["price"]["haiku"])), strict=True):
        label(te, f"A{r}", name)
        put(te, f"B{r}", pi, USD2)
        put(te, f"C{r}", po, USD2)
    header(te, 8, ["Agent (average tokens per call)", "Input tokens", "Output tokens", "Cost on Sonnet 5",
                   "Cost on Haiku 4.5"])
    for k, (name, i_tok, o_tok) in enumerate(I["agents"]):
        r = 9 + k
        label(te, f"A{r}", name)
        put(te, f"B{r}", i_tok, NUM)
        put(te, f"C{r}", o_tok, NUM)
        put(te, f"D{r}", f"=(B{r}*$B$5+C{r}*$C$5)/1000000", USD4)
        put(te, f"E{r}", f"=(B{r}*$B$6+C{r}*$C$6)/1000000", USD4)
    label(te, "A13", "Share of P1 incidents")
    put(te, "B13", I["p1_share"], PCT)
    label(te, "A14", "Semantic cache hit rate (triage + planner)")
    put(te, "B14", I["cache_hit"], PCT, note="Blueprint target 15-25 %. Demo re-runs hit 100 %; not a production figure.")
    label(te, "A15", "Overhead multiplier (re-plans, validation retries)")
    put(te, "B15", I["overhead"], '0.00"x"')

    header(te, 17, ["Token optimization strategy", "P1 $/incident", "P2 $/incident", "Blended",
                    "After cache", "With overhead"])
    label(te, "A18", "All Sonnet 5 (no routing, no cache)")
    put(te, "B18", "=SUM(D9:D11)", USD4)
    put(te, "C18", "=SUM(D9:D11)", USD4)
    label(te, "A19", "Tiered routing (Sonnet for P1, Haiku otherwise)")
    put(te, "B19", "=D9+D10+E11", USD4)
    put(te, "C19", "=SUM(E9:E11)", USD4)
    label(te, "A20", "Tiered routing + semantic cache (as built)", bold=True)
    put(te, "B20", "=B19", USD4)
    put(te, "C20", "=C19", USD4)
    for r in (18, 19, 20):
        put(te, f"D{r}", f"=$B$13*B{r}+(1-$B$13)*C{r}", USD4)
    put(te, "E18", "=D18", USD4)
    put(te, "E19", "=D19", USD4)
    put(te, "E20", "=D20-$B$14*($B$13*(D9+D10)+(1-$B$13)*(E9+E10))", USD4)
    for r in (18, 19, 20):
        put(te, f"F{r}", f"=E{r}*$B$15", USD4, fill=KEY_FILL if r == 20 else None)
    label(te, "A22", "Saving vs all-Sonnet")
    put(te, "F22", "=1-F20/F18", PCT)

    header(te, 24, ["3-year LLM production cost by strategy", *Y, "3-year total"])
    for k, r in enumerate((25, 26, 27)):
        label(te, f"A{r}", te[f"A{18 + k}"].value)
        for i, col in enumerate("BCD"):
            put(te, f"{col}{r}", f"=TCO!{col}4*12*$F${18 + k}", USD)
        put(te, f"E{r}", f"=SUM(B{r}:D{r})", USD)

    # ---------------- TCO
    tc = wb.create_sheet("TCO")
    tc["A1"].value, tc["A1"].font = "Total cost of ownership", TITLE
    header(tc, 3, ["Line ($ / month unless noted)", *Y], [46, 14, 14, 14])
    label(tc, "A4", "Incidents per month")
    put(tc, "B4", "=Inputs!B5", NUM)
    put(tc, "C4", "=B4*(1+Inputs!$B$6)", NUM)
    put(tc, "D4", "=C4*(1+Inputs!$B$6)", NUM)
    lines = [(5, "Infrastructure subtotal", "=SUM(Inputs!{c}28:{c}35)"),
             (6, "LLM production (all incidents x cost per incident)", "={c}4*'Token Economics'!$F$20"),
             (7, "LLM evaluation + staging", "=Inputs!$B$17"),
             (8, "Platform maintenance", "=Inputs!{c}16*Inputs!$B$4")]
    for r, text, f in lines:
        label(tc, f"A{r}", text)
        for col in "BCD":
            put(tc, f"{col}{r}", f.format(c=col), USD)
    label(tc, "A9", "Total per month", bold=True)
    label(tc, "A10", "Total per year", bold=True)
    for col in "BCD":
        put(tc, f"{col}9", f"=SUM({col}5:{col}8)", USD, BOLD)
        put(tc, f"{col}10", f"={col}9*12", USD, BOLD)
    tc["A12"].value, tc["A12"].font = "Year 0 (one-time)", BOLD
    label(tc, "A13", "Capstone build (hours x rate)")
    put(tc, "B13", "=Inputs!B22*Inputs!B4", USD)
    label(tc, "A14", "Productionization")
    put(tc, "B14", "=Inputs!B23", USD)
    label(tc, "A15", "Landing zone + hardening")
    put(tc, "B15", "=Inputs!B24", USD)
    label(tc, "A16", "Policy authoring + review")
    put(tc, "B16", "=Inputs!B25", USD)
    label(tc, "A17", "Total Year 0", bold=True)
    put(tc, "B17", "=SUM(B13:B16)", USD, BOLD)

    # ---------------- Benefits
    be = wb.create_sheet("Benefits")
    be["A1"].value, be["A1"].font = "Benefits: SRE capacity freed (valued at the loaded rate, not headcount cuts)", TITLE
    header(be, 3, ["Line (per month unless noted)", *Y], [46, 14, 14, 14])
    brows = [
        (4, "Incidents through the platform", "=TCO!{c}4*Inputs!{c}7", NUM1),
        (5, "Automated incidents", "={c}4*Inputs!{c}8", NUM1),
        (6, "HITL incidents", "={c}4*Inputs!{c}9", NUM1),
        (7, "Escalated incidents", "={c}4*Inputs!{c}10", NUM1),
        (8, "Minutes saved: automated", "={c}5*Inputs!$B$11", NUM),
        (9, "Minutes saved: HITL", "={c}6*Inputs!$B$12", NUM),
        (10, "Minutes saved: escalated", "={c}7*Inputs!$B$13", NUM),
        (11, "Bad-remediation cleanup (minutes)", "=-({c}5+{c}6)*Inputs!{c}14*Inputs!$B$15", NUM),
        (12, "Net minutes saved", "=SUM({c}8:{c}11)", NUM),
        (13, "SRE hours freed", "={c}12/60", NUM1),
        (14, "Annual value ($ / year)", "={c}13*12*Inputs!$B$4", USD),
    ]
    for r, text, f, fmt in brows:
        label(be, f"A{r}", text, bold=r in (12, 14))
        for col in "BCD":
            put(be, f"{col}{r}", f.format(c=col), fmt, BOLD if r == 14 else None)

    # ---------------- Summary
    su = wb.create_sheet("Summary", 0)
    su["A1"].value, su["A1"].font = "CloudScale incident platform: 3-year TCO / ROI", TITLE
    su["A2"].value = "All figures are formulas over the Inputs and Token Economics sheets."
    header(su, 4, ["Year", "Platform cost", "Benefit", "Net", "Cumulative", "Discount factor", "Present value"],
           [22, 16, 16, 16, 16, 16, 16])
    for k, (yr, cost, ben) in enumerate((("Y0", "=TCO!B17", "=0"), ("Y1", "=TCO!B10", "=Benefits!B14"),
                                         ("Y2", "=TCO!C10", "=Benefits!C14"), ("Y3", "=TCO!D10", "=Benefits!D14"))):
        r = 5 + k
        su[f"A{r}"].value, su[f"A{r}"].font = yr, NORMAL
        put(su, f"B{r}", cost, USD)
        put(su, f"C{r}", ben, USD)
        put(su, f"D{r}", f"=C{r}-B{r}", USD)
        put(su, f"E{r}", "=D5" if r == 5 else f"=E{r - 1}+D{r}", USD)
        put(su, f"F{r}", f"=1/(1+Inputs!$B$18)^{k}", "0.0000")
        put(su, f"G{r}", f"=D{r}*F{r}", USD)
    label(su, "A9", "3-year total", bold=True)
    for col in "BCDG":
        put(su, f"{col}9", f"=SUM({col}5:{col}8)", USD, BOLD)
    kpis = [
        (11, "NPV (10 %)", "=G9", USD),
        (12, "3-year net", "=E8", USD),
        (13, "Benefit-cost ratio", "=IF(B9=0,0,C9/B9)", '0.00"x"'),
        (14, "Net ROI", "=IF(B9=0,0,E8/B9)", PCT),
        (15, "Payback (month of operation)",
         '=IF(E6>=0,ROUNDUP(-E5/(D6/12),0),IF(E7>=0,12+ROUNDUP(-E6/(D7/12),0),IF(E8>=0,24+ROUNDUP(-E7/(D8/12),0),"> 36")))',
         "0"),
        (16, "LLM share of 3-year cost", "=(SUM(TCO!B6:D7)*12)/B9", PCT),
        (17, "LLM cost per incident (as built)", "='Token Economics'!F20", USD4),
        (18, "LLM cost per SRE-minute saved (automated path)", "=B17/Inputs!B11", USD4),
        (19, "SRE hours freed, Y1 / Y3", '=TEXT(Benefits!B13*12,"#,##0")&" h / "&TEXT(Benefits!D13*12,"#,##0")&" h"',
         None),
    ]
    for r, text, f, fmt in kpis:
        label(su, f"A{r}", text, bold=True)
        put(su, f"B{r}", f, fmt, BOLD, fill=KEY_FILL if r in (11, 12, 15) else None)

    # ---------------- Sensitivity
    se = wb.create_sheet("Sensitivity")
    se["A1"].value, se["A1"].font = "Sensitivity: 3-year net", TITLE
    header(se, 3, ["Case", "3-year net", "Change vs base"], [48, 16, 16])
    cases = [
        ("Base", "=Summary!E8"),
        ("Lower automation (points automated -> HITL, Inputs!B39)",
         "=B4-SUMPRODUCT(Benefits!B4:D4)*Inputs!B39*(Inputs!B11-Inputs!B12)*12/60*Inputs!B4"),
        ("Slower adoption (Inputs row 40)",
         ("=B4-(Benefits!B14*(1-Inputs!B40/Inputs!B7)+Benefits!C14*(1-Inputs!C40/Inputs!C7)"
          "+Benefits!D14*(1-Inputs!D40/Inputs!D7))")),
        ("Lower volume (Inputs!B41 fewer incidents)",
         "=B4-Inputs!B41*Summary!C9+Inputs!B41*SUM(TCO!B6:D6)*12"),
        ("LLM prices x2", "=B4-SUM(TCO!B6:D7)*12"),
        ("Productionization +50 %", "=B4-0.5*Inputs!B23"),
        ("Azure warm DR", "=B4-Inputs!B19*36"),
    ]
    for k, (text, f) in enumerate(cases):
        r = 4 + k
        label(se, f"A{r}", text)
        put(se, f"B{r}", f, USD)
        put(se, f"C{r}", f"=B{r}-$B$4", USD)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "B4" if sheet.title != "Summary" else "B5"
    wb.calculation.fullCalcOnLoad = True     # Excel recalculates every formula on open
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def verify(expected: dict) -> None:
    import formulas
    xl = formulas.ExcelModel().loads(str(OUT)).finish()
    sol = xl.calculate()

    def val(sheet, ref):
        for k, v in sol.items():
            if k.upper().endswith(f"{sheet.upper()}'!{ref}") or k.upper().endswith(f"]{sheet.upper()}!{ref}"):
                return float(v.value[0, 0])
        raise KeyError(f"{sheet}!{ref}")

    checks = {
        "cost per incident": (val("TOKEN ECONOMICS", "F20"), expected["cost_per_incident"]),
        "Year 0": (val("TCO", "B17"), expected["capex"]),
        "OPEX Y1": (val("TCO", "B10"), expected["opex"][0]),
        "Benefit Y3": (val("BENEFITS", "D14"), expected["benefit"][2]),
        "NPV": (val("SUMMARY", "B11"), expected["npv"]),
        "3-year net": (val("SUMMARY", "B12"), expected["cum"][3]),
        "payback month": (val("SUMMARY", "B15"), expected["payback_month"]),
        "lower automation": (val("SENSITIVITY", "B5"),
                             expected["sensitivity"]["Lower automation (15 pts automated -> HITL)"]),
        "lower volume": (val("SENSITIVITY", "B7"), expected["sensitivity"]["Lower volume (-30 % incidents)"]),
    }
    bad = {k: v for k, v in checks.items() if abs(v[0] - v[1]) > 0.01 * max(1, abs(v[1]) / 1000)}
    for k, (x, p) in checks.items():
        print(f"  {'OK ' if k not in bad else 'BAD'} {k:18} workbook={x:,.4f}  python={p:,.4f}")
    if bad:
        raise SystemExit(f"{len(bad)} mismatch(es)")
    print("workbook formulas match the independent Python computation")


if __name__ == "__main__":
    results = compute()
    build()
    print(f"wrote {OUT}")
    print(json.dumps({k: v for k, v in results.items() if k not in ("agent_cost_sonnet", "agent_cost_haiku")},
                     indent=1, default=lambda x: round(x, 4)))
    if "--verify" in sys.argv:
        verify(results)
